from io import BytesIO
from datetime import datetime
from typing import Any, Dict, List

from docx import Document
from docx.shared import (
    Inches,
    Pt,
    RGBColor
)
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH
)
from docx.enum.table import (
    WD_TABLE_ALIGNMENT,
    WD_CELL_VERTICAL_ALIGNMENT
)
from docx.enum.section import (
    WD_SECTION
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# =========================================================
# COLORS
# =========================================================

PRIMARY_COLOR = "2563EB"
SECONDARY_COLOR = "1E40AF"
SUCCESS_COLOR = "16A34A"
WARNING_COLOR = "D97706"
DANGER_COLOR = "DC2626"
LIGHT_BLUE = "EFF6FF"
LIGHT_GREEN = "F0FDF4"
LIGHT_ORANGE = "FFF7ED"
LIGHT_RED = "FEF2F2"
LIGHT_GRAY = "F3F4F6"
DARK_GRAY = "374151"


# =========================================================
# HELPER FUNCTIONS
# =========================================================

def ensure_list(value: Any) -> List[Any]:
    """
    Convert a value into a list.
    """

    if value is None:
        return []

    if isinstance(value, list):
        return value

    if isinstance(value, tuple):
        return list(value)

    if isinstance(value, str):

        if not value.strip():
            return []

        return [
            item.strip()
            for item in value.split(",")
            if item.strip()
        ]

    return [value]


def safe_int(
    value: Any,
    default: int = 0
) -> int:
    """
    Safely convert a value into an integer.
    """

    try:
        return int(float(value))
    except (
        ValueError,
        TypeError
    ):
        return default


def safe_float(
    value: Any,
    default: float = 0.0
) -> float:
    """
    Safely convert a value into a float.
    """

    try:
        return float(value)
    except (
        ValueError,
        TypeError
    ):
        return default


def clean_text(value: Any) -> str:
    """
    Convert any value into readable text.
    """

    if value is None:
        return ""

    if isinstance(value, dict):

        return ", ".join(
            f"{key}: {val}"
            for key, val in value.items()
        )

    if isinstance(value, list):

        return ", ".join(
            str(item)
            for item in value
        )

    return str(value)


# =========================================================
# DOCUMENT STYLING
# =========================================================

def set_cell_background(
    cell,
    color: str
):
    """
    Set background color of a table cell.
    """

    tc_pr = cell._tc.get_or_add_tcPr()

    shd = tc_pr.find(
        qn("w:shd")
    )

    if shd is None:

        shd = OxmlElement(
            "w:shd"
        )

        tc_pr.append(
            shd
        )

    shd.set(
        qn("w:fill"),
        color
    )


def set_cell_text_color(
    cell,
    color: str
):
    """
    Set text color inside a table cell.
    """

    for paragraph in cell.paragraphs:

        for run in paragraph.runs:

            run.font.color.rgb = RGBColor.from_string(
                color
            )


def set_cell_margins(
    cell,
    top: int = 100,
    start: int = 100,
    bottom: int = 100,
    end: int = 100
):
    """
    Set cell margins.
    """

    tc = cell._tc

    tc_pr = tc.get_or_add_tcPr()

    tc_mar = tc_pr.first_child_found_in(
        "w:tcMar"
    )

    if tc_mar is None:

        tc_mar = OxmlElement(
            "w:tcMar"
        )

        tc_pr.append(
            tc_mar
        )

    for margin, value in [
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end)
    ]:

        node = tc_mar.find(
            qn(f"w:{margin}")
        )

        if node is None:

            node = OxmlElement(
                f"w:{margin}"
            )

            tc_mar.append(
                node
            )

        node.set(
            qn("w:w"),
            str(value)
        )

        node.set(
            qn("w:type"),
            "dxa"
        )


def set_column_width(
    cell,
    width_inches: float
):
    """
    Set approximate table column width.
    """

    cell.width = Inches(
        width_inches
    )


def add_page_number(
    paragraph
):
    """
    Add page number field.
    """

    run = paragraph.add_run()

    fld_char1 = OxmlElement(
        "w:fldChar"
    )

    fld_char1.set(
        qn("w:fldCharType"),
        "begin"
    )

    instr_text = OxmlElement(
        "w:instrText"
    )

    instr_text.set(
        qn("xml:space"),
        "preserve"
    )

    instr_text.text = "PAGE"

    fld_char2 = OxmlElement(
        "w:fldChar"
    )

    fld_char2.set(
        qn("w:fldCharType"),
        "end"
    )

    run._r.append(
        fld_char1
    )

    run._r.append(
        instr_text
    )

    run._r.append(
        fld_char2
    )


# =========================================================
# DOCUMENT SETUP
# =========================================================

def setup_document(
    document: Document
):
    """
    Configure document margins and default fonts.
    """

    section = document.sections[0]

    section.top_margin = Inches(
        0.7
    )

    section.bottom_margin = Inches(
        0.7
    )

    section.left_margin = Inches(
        0.7
    )

    section.right_margin = Inches(
        0.7
    )

    styles = document.styles

    normal = styles["Normal"]

    normal.font.name = "Aptos"

    normal.font.size = Pt(
        10
    )

    # Heading 1

    heading1 = styles["Heading 1"]

    heading1.font.name = "Aptos Display"

    heading1.font.size = Pt(
        18
    )

    heading1.font.bold = True

    heading1.font.color.rgb = RGBColor.from_string(
        PRIMARY_COLOR
    )

    # Heading 2

    heading2 = styles["Heading 2"]

    heading2.font.name = "Aptos Display"

    heading2.font.size = Pt(
        14
    )

    heading2.font.bold = True

    heading2.font.color.rgb = RGBColor.from_string(
        SECONDARY_COLOR
    )


# =========================================================
# TITLE
# =========================================================

def add_title_page(
    document: Document,
    candidate_name: str = "Candidate",
    target_role: str = ""
):
    """
    Add professional report title.
    """

    paragraph = document.add_paragraph()

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    paragraph.space_before = Pt(
        80
    )

    run = paragraph.add_run(
        "AI RESUME &\n"
        "SKILL GAP REPORT"
    )

    run.font.name = "Aptos Display"

    run.font.size = Pt(
        28
    )

    run.font.bold = True

    run.font.color.rgb = RGBColor.from_string(
        PRIMARY_COLOR
    )

    # Candidate name

    paragraph = document.add_paragraph()

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = paragraph.add_run(
        candidate_name
    )

    run.font.size = Pt(
        18
    )

    run.font.bold = True

    run.font.color.rgb = RGBColor.from_string(
        DARK_GRAY
    )

    # Target role

    if target_role:

        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run(
            f"Target Role: {target_role}"
        )

        run.font.size = Pt(
            12
        )

        run.font.color.rgb = RGBColor.from_string(
            SECONDARY_COLOR
        )

    # Generated date

    paragraph = document.add_paragraph()

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = paragraph.add_run(
        datetime.now().strftime(
            "%d %B %Y"
        )
    )

    run.font.size = Pt(
        10
    )

    run.font.color.rgb = RGBColor.from_string(
        DARK_GRAY
    )

    # AI badge

    paragraph = document.add_paragraph()

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = paragraph.add_run(
        "Generated by AI Resume & Skill Gap Agent"
    )

    run.font.size = Pt(
        9
    )

    run.italic = True

    run.font.color.rgb = RGBColor.from_string(
        "6B7280"
    )

    document.add_page_break()


# =========================================================
# SECTION HEADING
# =========================================================

def add_section_heading(
    document: Document,
    title: str
):
    """
    Add styled section heading.
    """

    paragraph = document.add_paragraph(
        style="Heading 1"
    )

    paragraph.paragraph_format.space_before = Pt(
        12
    )

    paragraph.paragraph_format.space_after = Pt(
        8
    )

    run = paragraph.add_run(
        title
    )

    return paragraph


# =========================================================
# SUBHEADING
# =========================================================

def add_subheading(
    document: Document,
    title: str
):
    """
    Add styled subsection heading.
    """

    paragraph = document.add_paragraph(
        style="Heading 2"
    )

    paragraph.paragraph_format.space_before = Pt(
        8
    )

    paragraph.paragraph_format.space_after = Pt(
        5
    )

    paragraph.add_run(
        title
    )

    return paragraph


# =========================================================
# BULLET LIST
# =========================================================

def add_bullet_list(
    document: Document,
    items: Any,
    color: str = DARK_GRAY
):
    """
    Add bullet points.
    """

    items = ensure_list(
        items
    )

    for item in items:

        if isinstance(
            item,
            dict
        ):

            text = clean_text(
                item
            )

        else:

            text = clean_text(
                item
            )

        if not text:
            continue

        paragraph = document.add_paragraph(
            style="List Bullet"
        )

        paragraph.paragraph_format.space_after = Pt(
            2
        )

        run = paragraph.add_run(
            text
        )

        run.font.color.rgb = RGBColor.from_string(
            color
        )


# =========================================================
# SCORE BOX
# =========================================================

def add_score_box(
    document: Document,
    title: str,
    score: float,
    maximum: int = 100
):
    """
    Add a visual score box.
    """

    score = max(
        0,
        min(
            maximum,
            score
        )
    )

    percentage = (
        score / maximum
    ) * 100

    if percentage >= 80:

        color = SUCCESS_COLOR

        background = LIGHT_GREEN

    elif percentage >= 60:

        color = WARNING_COLOR

        background = LIGHT_ORANGE

    else:

        color = DANGER_COLOR

        background = LIGHT_RED

    table = document.add_table(
        rows=1,
        cols=2
    )

    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    table.autofit = True

    left = table.cell(
        0,
        0
    )

    right = table.cell(
        0,
        1
    )

    set_cell_background(
        left,
        background
    )

    set_cell_background(
        right,
        background
    )

    set_cell_margins(
        left
    )

    set_cell_margins(
        right
    )

    left_paragraph = left.paragraphs[0]

    left_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = left_paragraph.add_run(
        title
    )

    run.font.bold = True

    run.font.size = Pt(
        11
    )

    run.font.color.rgb = RGBColor.from_string(
        DARK_GRAY
    )

    right_paragraph = right.paragraphs[0]

    right_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = right_paragraph.add_run(
        f"{score:.0f}/{maximum}"
    )

    run.font.bold = True

    run.font.size = Pt(
        22
    )

    run.font.color.rgb = RGBColor.from_string(
        color
    )

    document.add_paragraph()


# =========================================================
# KEY-VALUE TABLE
# =========================================================

def add_key_value_table(
    document: Document,
    data: Dict[str, Any]
):
    """
    Add a two-column information table.
    """

    table = document.add_table(
        rows=0,
        cols=2
    )

    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    table.style = "Table Grid"

    for key, value in data.items():

        row = table.add_row()

        key_cell = row.cells[0]

        value_cell = row.cells[1]

        set_cell_background(
            key_cell,
            LIGHT_BLUE
        )

        set_cell_margins(
            key_cell
        )

        set_cell_margins(
            value_cell
        )

        key_cell.vertical_alignment = (
            WD_CELL_VERTICAL_ALIGNMENT.CENTER
        )

        value_cell.vertical_alignment = (
            WD_CELL_VERTICAL_ALIGNMENT.CENTER
        )

        key_paragraph = key_cell.paragraphs[0]

        key_run = key_paragraph.add_run(
            str(key)
        )

        key_run.bold = True

        key_run.font.color.rgb = RGBColor.from_string(
            SECONDARY_COLOR
        )

        value_paragraph = value_cell.paragraphs[0]

        value_paragraph.add_run(
            clean_text(value)
        )

    document.add_paragraph()


# =========================================================
# SKILLS TABLE
# =========================================================

def add_skills_table(
    document: Document,
    matched: Any,
    partial: Any,
    missing: Any
):
    """
    Create matched/partial/missing skills table.
    """

    table = document.add_table(
        rows=1,
        cols=3
    )

    table.style = "Table Grid"

    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    headers = [
        (
            "Matched Skills",
            SUCCESS_COLOR,
            LIGHT_GREEN
        ),
        (
            "Partial Skills",
            WARNING_COLOR,
            LIGHT_ORANGE
        ),
        (
            "Missing Skills",
            DANGER_COLOR,
            LIGHT_RED
        )
    ]

    for index, (
        title,
        color,
        background
    ) in enumerate(headers):

        cell = table.cell(
            0,
            index
        )

        set_cell_background(
            cell,
            background
        )

        set_cell_margins(
            cell
        )

        paragraph = cell.paragraphs[0]

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run(
            title
        )

        run.bold = True

        run.font.color.rgb = RGBColor.from_string(
            color
        )

    columns = [
        ensure_list(matched),
        ensure_list(partial),
        ensure_list(missing)
    ]

    max_rows = max(
        [
            len(column)
            for column in columns
        ] + [1]
    )

    for row_index in range(
        max_rows
    ):

        row = table.add_row()

        for column_index, column in enumerate(
            columns
        ):

            cell = row.cells[
                column_index
            ]

            set_cell_margins(
                cell
            )

            if row_index < len(
                column
            ):

                paragraph = cell.paragraphs[0]

                paragraph.add_run(
                    str(
                        column[
                            row_index
                        ]
                    )
                )

            else:

                cell.paragraphs[0].add_run(
                    "-"
                )

    document.add_paragraph()


# =========================================================
# RESUME SUMMARY
# =========================================================

def add_resume_summary(
    document: Document,
    resume_analysis: Dict[str, Any]
):
    """
    Add resume analysis section.
    """

    add_section_heading(
        document,
        "1. Resume Analysis"
    )

    data = resume_analysis.get(
        "data",
        resume_analysis
    )

    summary = data.get(
        "candidate_summary",
        ""
    )

    current_role = data.get(
        "current_role",
        ""
    )

    experience_level = data.get(
        "experience_level",
        ""
    )

    if summary:

        add_subheading(
            document,
            "Candidate Summary"
        )

        paragraph = document.add_paragraph(
            summary
        )

        paragraph.paragraph_format.space_after = Pt(
            8
        )

    add_key_value_table(
        document,
        {
            "Current Role": current_role or "Not specified",
            "Experience Level": (
                experience_level
                or "Not specified"
            )
        }
    )

    ats_score = safe_float(
        data.get(
            "ats_score",
            0
        )
    )

    add_score_box(
        document,
        "ATS Resume Score",
        ats_score
    )


# =========================================================
# STRENGTHS AND WEAKNESSES
# =========================================================

def add_strengths_weaknesses(
    document: Document,
    resume_analysis: Dict[str, Any]
):
    """
    Add strengths and weaknesses.
    """

    data = resume_analysis.get(
        "data",
        resume_analysis
    )

    add_subheading(
        document,
        "Strengths"
    )

    add_bullet_list(
        document,
        data.get(
            "strengths",
            []
        ),
        SUCCESS_COLOR
    )

    add_subheading(
        document,
        "Weaknesses"
    )

    add_bullet_list(
        document,
        data.get(
            "weaknesses",
            []
        ),
        DANGER_COLOR
    )

    add_subheading(
        document,
        "Missing Information"
    )

    add_bullet_list(
        document,
        data.get(
            "missing_information",
            []
        ),
        WARNING_COLOR
    )


# =========================================================
# SKILLS SECTION
# =========================================================

def add_skills_section(
    document: Document,
    skills: Dict[str, Any]
):
    """
    Add extracted skills.
    """

    add_section_heading(
        document,
        "2. Skills Profile"
    )

    data = skills.get(
        "data",
        skills
    )

    categories = [
        (
            "Technical Skills",
            "technical_skills"
        ),
        (
            "Programming Languages",
            "programming_languages"
        ),
        (
            "Frameworks",
            "frameworks"
        ),
        (
            "Databases",
            "databases"
        ),
        (
            "Cloud & Tools",
            "cloud_tools"
        ),
        (
            "Data Tools",
            "data_tools"
        ),
        (
            "AI / ML Tools",
            "ai_ml_tools"
        ),
        (
            "Soft Skills",
            "soft_skills"
        )
    ]

    for title, key in categories:

        values = ensure_list(
            data.get(
                key,
                []
            )
        )

        if not values:
            continue

        add_subheading(
            document,
            title
        )

        add_bullet_list(
            document,
            values
        )


# =========================================================
# JOB SECTION
# =========================================================

def add_job_section(
    document: Document,
    job_analysis: Dict[str, Any],
    job_title: str = "",
    company_name: str = ""
):
    """
    Add target job analysis.
    """

    add_section_heading(
        document,
        "3. Target Job Analysis"
    )

    data = job_analysis.get(
        "data",
        job_analysis
    )

    title = (
        data.get(
            "job_title",
            ""
        )
        or job_title
        or "Target Role"
    )

    add_key_value_table(
        document,
        {
            "Job Title": title,
            "Company": company_name or "Not specified"
        }
    )

    categories = [
        (
            "Required Skills",
            "required_skills"
        ),
        (
            "Preferred Skills",
            "preferred_skills"
        ),
        (
            "Programming Languages",
            "programming_languages"
        ),
        (
            "Frameworks",
            "frameworks"
        ),
        (
            "Databases",
            "databases"
        ),
        (
            "AI / ML Tools",
            "ai_ml_tools"
        ),
        (
            "Experience Requirements",
            "experience_requirements"
        )
    ]

    for title, key in categories:

        values = ensure_list(
            data.get(
                key,
                []
            )
        )

        if not values:
            continue

        add_subheading(
            document,
            title
        )

        add_bullet_list(
            document,
            values
        )


# =========================================================
# SKILL GAP SECTION
# =========================================================

def add_skill_gap_section(
    document: Document,
    skill_gap: Dict[str, Any]
):
    """
    Add skill-gap analysis.
    """

    add_section_heading(
        document,
        "4. Skill Gap Analysis"
    )

    data = skill_gap.get(
        "data",
        skill_gap
    )

    match_percentage = safe_float(
        data.get(
            "match_percentage",
            0
        )
    )

    add_score_box(
        document,
        "Job Match",
        match_percentage
    )

    add_skills_table(
        document,
        data.get(
            "matched_skills",
            []
        ),
        data.get(
            "partial_skills",
            []
        ),
        data.get(
            "missing_skills",
            []
        )
    )

    add_subheading(
        document,
        "Priority Skill Gaps"
    )

    add_bullet_list(
        document,
        data.get(
            "priority_gaps",
            []
        ),
        DANGER_COLOR
    )

    add_subheading(
        document,
        "Experience Gaps"
    )

    add_bullet_list(
        document,
        data.get(
            "experience_gaps",
            []
        ),
        WARNING_COLOR
    )

    add_subheading(
        document,
        "Recommendations"
    )

    add_bullet_list(
        document,
        data.get(
            "recommendations",
            []
        ),
        PRIMARY_COLOR
    )


# =========================================================
# RESUME IMPROVEMENTS
# =========================================================

def add_resume_improvements(
    document: Document,
    improvements: Dict[str, Any]
):
    """
    Add resume improvement recommendations.
    """

    add_section_heading(
        document,
        "5. Resume Improvement Plan"
    )

    data = improvements.get(
        "data",
        improvements
    )

    score = safe_float(
        data.get(
            "overall_score",
            0
        )
    )

    add_score_box(
        document,
        "Resume Quality Score",
        score
    )

    categories = [
        (
            "Summary Improvements",
            "summary_improvements"
        ),
        (
            "Experience Improvements",
            "experience_improvements"
        ),
        (
            "Skills Improvements",
            "skills_improvements"
        ),
        (
            "Project Improvements",
            "project_improvements"
        ),
        (
            "Education Improvements",
            "education_improvements"
        ),
        (
            "ATS Keywords",
            "ats_keywords_to_consider"
        ),
        (
            "Formatting Recommendations",
            "formatting_recommendations"
        ),
        (
            "Action Items",
            "action_items"
        )
    ]

    for title, key in categories:

        values = ensure_list(
            data.get(
                key,
                []
            )
        )

        if not values:
            continue

        add_subheading(
            document,
            title
        )

        add_bullet_list(
            document,
            values
        )


# =========================================================
# CAREER RECOMMENDATIONS
# =========================================================

def add_career_recommendations(
    document: Document,
    careers: Dict[str, Any]
):
    """
    Add career recommendations.
    """

    add_section_heading(
        document,
        "6. Career Recommendations"
    )

    data = careers.get(
        "data",
        careers
    )

    current_profile = data.get(
        "current_profile",
        ""
    )

    if current_profile:

        add_subheading(
            document,
            "Current Career Profile"
        )

        document.add_paragraph(
            current_profile
        )

    recommendations = ensure_list(
        data.get(
            "recommended_roles",
            []
        )
    )

    for index, role in enumerate(
        recommendations,
        start=1
    ):

        add_subheading(
            document,
            f"{index}. {clean_text(role.get('role', 'Recommended Role'))}"
            if isinstance(role, dict)
            else f"{index}. {role}"
        )

        if isinstance(
            role,
            dict
        ):

            fit = safe_float(
                role.get(
                    "fit_percentage",
                    0
                )
            )

            if fit:

                paragraph = document.add_paragraph()

                run = paragraph.add_run(
                    f"Role Fit: {fit:.0f}%"
                )

                run.bold = True

                run.font.color.rgb = RGBColor.from_string(
                    PRIMARY_COLOR
                )

            reason = role.get(
                "reason",
                ""
            )

            if reason:

                document.add_paragraph(
                    reason
                )

            required_skills = role.get(
                "required_skills",
                []
            )

            if required_skills:

                paragraph = document.add_paragraph()

                run = paragraph.add_run(
                    "Required Skills:"
                )

                run.bold = True

                add_bullet_list(
                    document,
                    required_skills
                )

            skill_gaps = role.get(
                "skill_gaps",
                []
            )

            if skill_gaps:

                paragraph = document.add_paragraph()

                run = paragraph.add_run(
                    "Skill Gaps:"
                )

                run.bold = True

                add_bullet_list(
                    document,
                    skill_gaps,
                    DANGER_COLOR
                )


# =========================================================
# LEARNING ROADMAP
# =========================================================

def add_learning_roadmap(
    document: Document,
    roadmap: Dict[str, Any]
):
    """
    Add learning roadmap.
    """

    add_section_heading(
        document,
        "7. Personalized Learning Roadmap"
    )

    data = roadmap.get(
        "data",
        roadmap
    )

    target_role = data.get(
        "target_role",
        ""
    )

    experience_level = data.get(
        "experience_level",
        ""
    )

    duration = data.get(
        "estimated_duration",
        ""
    )

    add_key_value_table(
        document,
        {
            "Target Role": target_role,
            "Experience Level": experience_level,
            "Estimated Duration": duration
        }
    )

    priority_order = ensure_list(
        data.get(
            "priority_order",
            []
        )
    )

    if priority_order:

        add_subheading(
            document,
            "Priority Order"
        )

        add_bullet_list(
            document,
            priority_order
        )

    phases = ensure_list(
        data.get(
            "phases",
            []
        )
    )

    for phase in phases:

        if not isinstance(
            phase,
            dict
        ):

            continue

        phase_number = phase.get(
            "phase",
            ""
        )

        title = phase.get(
            "title",
            "Learning Phase"
        )

        duration = phase.get(
            "duration",
            ""
        )

        heading = (
            f"Phase {phase_number}: {title}"
        )

        if duration:

            heading += (
                f" ({duration})"
            )

        add_subheading(
            document,
            heading
        )

        objectives = phase.get(
            "objectives",
            []
        )

        if objectives:

            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                "Objectives"
            )

            run.bold = True

            add_bullet_list(
                document,
                objectives
            )

        skills = phase.get(
            "skills",
            []
        )

        if skills:

            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                "Skills"
            )

            run.bold = True

            add_bullet_list(
                document,
                skills
            )

        projects = phase.get(
            "practice_projects",
            []
        )

        if projects:

            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                "Practice Projects"
            )

            run.bold = True

            add_bullet_list(
                document,
                projects
            )

        criteria = phase.get(
            "completion_criteria",
            []
        )

        if criteria:

            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                "Completion Criteria"
            )

            run.bold = True

            add_bullet_list(
                document,
                criteria,
                SUCCESS_COLOR
            )

    portfolio_projects = ensure_list(
        data.get(
            "portfolio_projects",
            []
        )
    )

    if portfolio_projects:

        add_subheading(
            document,
            "Portfolio Projects"
        )

        add_bullet_list(
            document,
            portfolio_projects,
            PRIMARY_COLOR
        )

    interview = ensure_list(
        data.get(
            "interview_preparation",
            []
        )
    )

    if interview:

        add_subheading(
            document,
            "Interview Preparation"
        )

        add_bullet_list(
            document,
            interview
        )


# =========================================================
# INTERVIEW SECTION
# =========================================================

def add_interview_questions(
    document: Document,
    interview: Dict[str, Any]
):
    """
    Add interview preparation questions.
    """

    add_section_heading(
        document,
        "8. Interview Preparation"
    )

    data = interview.get(
        "data",
        interview
    )

    categories = [
        (
            "Technical Questions",
            "technical_questions"
        ),
        (
            "Behavioral Questions",
            "behavioral_questions"
        ),
        (
            "Project Questions",
            "project_questions"
        ),
        (
            "Resume-Based Questions",
            "resume_based_questions"
        ),
        (
            "Questions to Practice",
            "questions_to_practice"
        )
    ]

    for title, key in categories:

        questions = ensure_list(
            data.get(
                key,
                []
            )
        )

        if not questions:
            continue

        add_subheading(
            document,
            title
        )

        for index, question in enumerate(
            questions,
            start=1
        ):

            paragraph = document.add_paragraph()

            paragraph.paragraph_format.space_after = Pt(
                4
            )

            run = paragraph.add_run(
                f"{index}. {clean_text(question)}"
            )

            run.font.size = Pt(
                10
            )


# =========================================================
# FOOTER
# =========================================================

def add_footer(
    document: Document
):
    """
    Add footer to every page.
    """

    for section in document.sections:

        footer = section.footer

        paragraph = footer.paragraphs[0]

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run(
            "AI Resume & Skill Gap Agent  •  "
        )

        run.font.size = Pt(
            8
        )

        run.font.color.rgb = RGBColor.from_string(
            "6B7280"
        )

        add_page_number(
            paragraph
        )


# =========================================================
# MAIN REPORT GENERATOR
# =========================================================

class ReportGenerator:
    """
    Creates professional DOCX reports.
    """

    def __init__(
        self,
        title: str = "AI Resume & Skill Gap Report"
    ):

        self.title = title


    # =====================================================
    # CREATE DOCUMENT
    # =====================================================

    def create_document(
        self,
        analysis: Dict[str, Any],
        candidate_name: str = "Candidate",
        target_role: str = "",
        job_title: str = "",
        company_name: str = "",
        interview_questions: Dict[str, Any] = None
    ) -> Document:
        """
        Create a complete DOCX document.
        """

        document = Document()

        setup_document(
            document
        )

        # -------------------------------------------------
        # Determine target role
        # -------------------------------------------------

        final_target_role = (
            target_role
            or job_title
        )

        # -------------------------------------------------
        # Title
        # -------------------------------------------------

        add_title_page(
            document,
            candidate_name=candidate_name,
            target_role=final_target_role
        )

        # -------------------------------------------------
        # Resume analysis
        # -------------------------------------------------

        resume_analysis = analysis.get(
            "resume_analysis",
            {}
        )

        if resume_analysis:

            add_resume_summary(
                document,
                resume_analysis
            )

            add_strengths_weaknesses(
                document,
                resume_analysis
            )

        # -------------------------------------------------
        # Skills
        # -------------------------------------------------

        skills = analysis.get(
            "skills",
            {}
        )

        if skills:

            add_skills_section(
                document,
                skills
            )

        # -------------------------------------------------
        # Job analysis
        # -------------------------------------------------

        job_analysis = analysis.get(
            "job_analysis",
            {}
        )

        if job_analysis:

            add_job_section(
                document,
                job_analysis,
                job_title=job_title,
                company_name=company_name
            )

        # -------------------------------------------------
        # Skill gap
        # -------------------------------------------------

        skill_gap = analysis.get(
            "skill_gap",
            {}
        )

        if skill_gap:

            add_skill_gap_section(
                document,
                skill_gap
            )

        # -------------------------------------------------
        # Resume improvements
        # -------------------------------------------------

        improvements = analysis.get(
            "resume_improvements",
            {}
        )

        if improvements:

            add_resume_improvements(
                document,
                improvements
            )

        # -------------------------------------------------
        # Career recommendations
        # -------------------------------------------------

        careers = analysis.get(
            "career_recommendations",
            {}
        )

        if careers:

            add_career_recommendations(
                document,
                careers
            )

        # -------------------------------------------------
        # Learning roadmap
        # -------------------------------------------------

        roadmap = analysis.get(
            "roadmap",
            {}
        )

        if roadmap:

            add_learning_roadmap(
                document,
                roadmap
            )

        # -------------------------------------------------
        # Interview questions
        # -------------------------------------------------

        if interview_questions:

            add_interview_questions(
                document,
                interview_questions
            )

        # -------------------------------------------------
        # Final recommendations
        # -------------------------------------------------

        document.add_page_break()

        add_section_heading(
            document,
            "Final Action Plan"
        )

        action_items = []

        # Missing skills

        skill_gap_data = skill_gap.get(
            "data",
            {}
        )

        for skill in ensure_list(
            skill_gap_data.get(
                "missing_skills",
                []
            )
        ):

            action_items.append(
                f"Develop the skill: {skill}"
            )

        # Resume actions

        improvement_data = improvements.get(
            "data",
            {}
        )

        action_items.extend(
            ensure_list(
                improvement_data.get(
                    "action_items",
                    []
                )
            )
        )

        # Career actions

        career_data = careers.get(
            "data",
            {}
        )

        action_items.extend(
            ensure_list(
                career_data.get(
                    "next_steps",
                    []
                )
            )
        )

        # Roadmap actions

        roadmap_data = roadmap.get(
            "data",
            {}
        )

        action_items.extend(
            ensure_list(
                roadmap_data.get(
                    "final_recommendations",
                    []
                )
            )
        )

        # Remove duplicates

        unique_actions = []

        seen = set()

        for action in action_items:

            text = clean_text(
                action
            )

            normalized = text.lower()

            if (
                text
                and normalized not in seen
            ):

                seen.add(
                    normalized
                )

                unique_actions.append(
                    text
                )

        if unique_actions:

            add_bullet_list(
                document,
                unique_actions,
                PRIMARY_COLOR
            )

        else:

            document.add_paragraph(
                "Continue building relevant skills, "
                "improving the resume, and practicing "
                "for the target role."
            )

        # -------------------------------------------------
        # Footer
        # -------------------------------------------------

        add_footer(
            document
        )

        return document


    # =====================================================
    # GENERATE BYTES
    # =====================================================

    def generate_bytes(
        self,
        analysis: Dict[str, Any],
        candidate_name: str = "Candidate",
        target_role: str = "",
        job_title: str = "",
        company_name: str = "",
        interview_questions: Dict[str, Any] = None
    ) -> bytes:
        """
        Generate DOCX report as bytes.

        Useful for Streamlit download_button.
        """

        document = self.create_document(
            analysis=analysis,
            candidate_name=candidate_name,
            target_role=target_role,
            job_title=job_title,
            company_name=company_name,
            interview_questions=interview_questions
        )

        buffer = BytesIO()

        document.save(
            buffer
        )

        buffer.seek(
            0
        )

        return buffer.getvalue()


    # =====================================================
    # SAVE TO FILE
    # =====================================================

    def save_to_file(
        self,
        analysis: Dict[str, Any],
        output_path: str,
        candidate_name: str = "Candidate",
        target_role: str = "",
        job_title: str = "",
        company_name: str = "",
        interview_questions: Dict[str, Any] = None
    ) -> str:
        """
        Save report to a DOCX file.
        """

        document = self.create_document(
            analysis=analysis,
            candidate_name=candidate_name,
            target_role=target_role,
            job_title=job_title,
            company_name=company_name,
            interview_questions=interview_questions
        )

        document.save(
            output_path
        )

        return output_path


# =========================================================
# CONVENIENCE FUNCTION
# =========================================================

def generate_report(
    analysis: Dict[str, Any],
    candidate_name: str = "Candidate",
    target_role: str = "",
    job_title: str = "",
    company_name: str = "",
    interview_questions: Dict[str, Any] = None
) -> bytes:
    """
    Convenience function for generating a DOCX report.
    """

    generator = ReportGenerator()

    return generator.generate_bytes(
        analysis=analysis,
        candidate_name=candidate_name,
        target_role=target_role,
        job_title=job_title,
        company_name=company_name,
        interview_questions=interview_questions
    )


# =========================================================
# TEST
# =========================================================

if __name__ == "__main__":

    print(
        "=" * 70
    )

    print(
        "AI RESUME & SKILL GAP AGENT"
    )

    print(
        "Report Generator Test"
    )

    print(
        "=" * 70
    )

    sample_analysis = {

        "resume_analysis": {

            "data": {

                "candidate_summary":
                    "Python developer with "
                    "experience in data analysis.",

                "current_role":
                    "Python Developer",

                "experience_level":
                    "Junior",

                "ats_score":
                    72,

                "strengths": [
                    "Python programming",
                    "SQL knowledge",
                    "Data analysis"
                ],

                "weaknesses": [
                    "Limited machine learning experience"
                ],

                "missing_information": [
                    "Quantified achievements"
                ]
            }
        },

        "skills": {

            "data": {

                "technical_skills": [
                    "Python",
                    "SQL",
                    "Pandas"
                ],

                "programming_languages": [
                    "Python"
                ],

                "frameworks": [
                    "Streamlit"
                ],

                "databases": [
                    "SQLite"
                ],

                "data_tools": [
                    "Pandas"
                ]
            }
        },

        "skill_gap": {

            "data": {

                "match_percentage":
                    68,

                "matched_skills": [
                    "Python",
                    "SQL",
                    "Pandas"
                ],

                "partial_skills": [
                    "Machine Learning"
                ],

                "missing_skills": [
                    "Scikit-learn",
                    "Statistics"
                ],

                "priority_gaps": [
                    "Machine Learning",
                    "Statistics"
                ],

                "recommendations": [
                    "Learn supervised learning",
                    "Practice statistical analysis"
                ]
            }
        },

        "resume_improvements": {

            "data": {

                "overall_score":
                    72,

                "summary_improvements": [
                    "Add measurable achievements"
                ],

                "experience_improvements": [
                    "Use action-oriented bullet points"
                ],

                "ats_keywords_to_consider": [
                    "Machine Learning",
                    "Scikit-learn"
                ],

                "action_items": [
                    "Add measurable results to projects"
                ]
            }
        },

        "career_recommendations": {

            "data": {

                "recommended_roles": [

                    {
                        "role":
                            "Junior Data Analyst",

                        "fit_percentage":
                            85,

                        "reason":
                            "Strong Python and SQL foundation",

                        "required_skills": [
                            "Python",
                            "SQL"
                        ],

                        "skill_gaps": [
                            "Advanced statistics"
                        ]
                    }
                ],

                "next_steps": [
                    "Build a data-analysis portfolio"
                ]
            }
        },

        "roadmap": {

            "data": {

                "target_role":
                    "Data Scientist",

                "experience_level":
                    "beginner",

                "estimated_duration":
                    "4-6 months",

                "priority_order": [
                    "Statistics",
                    "Machine Learning",
                    "Scikit-learn"
                ],

                "phases": [

                    {
                        "phase":
                            1,

                        "title":
                            "Statistics Fundamentals",

                        "duration":
                            "4 weeks",

                        "skills": [
                            "Statistics",
                            "Probability"
                        ],

                        "objectives": [
                            "Understand descriptive statistics"
                        ],

                        "practice_projects": [
                            "Analyze a public dataset"
                        ],

                        "completion_criteria": [
                            "Complete one analysis project"
                        ]
                    }
                ],

                "portfolio_projects": [
                    "Build a customer churn prediction project"
                ],

                "interview_preparation": [
                    "Practice Python interview questions"
                ],

                "final_recommendations": [
                    "Build two portfolio projects"
                ]
            }
        }
    }

    generator = ReportGenerator()

    output_file = "sample_resume_skill_gap_report.docx"

    generator.save_to_file(
        analysis=sample_analysis,
        output_path=output_file,
        candidate_name="John Doe",
        target_role="Data Scientist"
    )

    print(
        f"\n✅ Report generated successfully:"
    )

    print(
        output_file
    )

    print(
        "\n" + "=" * 70
    )

