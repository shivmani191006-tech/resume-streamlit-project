import io
import re
from pathlib import Path
from typing import Any, Dict, Optional


# =========================================================
# OPTIONAL PDF IMPORT
# =========================================================

try:

    import pypdf as PyPDF2

    PDF_AVAILABLE = True

except ImportError:

    try:

        import PyPDF2

        PDF_AVAILABLE = True

    except ImportError:

        PDF_AVAILABLE = False


# =========================================================
# OPTIONAL DOCX IMPORT
# =========================================================

try:

    from docx import Document

    DOCX_AVAILABLE = True

except ImportError:

    DOCX_AVAILABLE = False


# =========================================================
# SUPPORTED FILE TYPES
# =========================================================

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt"
}


# =========================================================
# MAXIMUM TEXT SIZE
# =========================================================

MAX_TEXT_LENGTH = 100_000


# =========================================================
# FILE VALIDATION
# =========================================================

def validate_file(
    file_name: str
) -> Dict[str, Any]:
    """
    Validate uploaded resume file.

    Returns:
        Dictionary containing validation result.
    """

    if not file_name:

        return {
            "valid": False,
            "message": "No file name provided."
        }

    extension = Path(
        file_name
    ).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:

        return {
            "valid": False,
            "message": (
                f"Unsupported file type: {extension}. "
                f"Supported formats are PDF, DOCX and TXT."
            )
        }

    return {
        "valid": True,
        "extension": extension,
        "message": "File type is supported."
    }


# =========================================================
# CLEAN TEXT
# =========================================================

def clean_text(
    text: str
) -> str:
    """
    Clean extracted resume text.

    Operations:
        - Remove null characters
        - Normalize whitespace
        - Remove excessive blank lines
        - Remove repeated spaces
        - Preserve readable structure
    """

    if not text:

        return ""

    # Remove null characters

    text = text.replace(
        "\x00",
        " "
    )

    # Normalize Windows line endings

    text = text.replace(
        "\r\n",
        "\n"
    )

    text = text.replace(
        "\r",
        "\n"
    )

    # Replace tabs with spaces

    text = text.replace(
        "\t",
        " "
    )

    # Remove trailing spaces

    lines = []

    for line in text.split("\n"):

        line = re.sub(
            r"[ ]+",
            " ",
            line
        )

        line = line.strip()

        lines.append(
            line
        )

    # Remove excessive blank lines

    cleaned_lines = []

    previous_blank = False

    for line in lines:

        if not line:

            if previous_blank:

                continue

            previous_blank = True

        else:

            previous_blank = False

        cleaned_lines.append(
            line
        )

    text = "\n".join(
        cleaned_lines
    )

    # Remove excessive spaces around newlines

    text = re.sub(
        r"\n[ ]+",
        "\n",
        text
    )

    text = re.sub(
        r"[ ]+\n",
        "\n",
        text
    )

    # Limit text size

    if len(text) > MAX_TEXT_LENGTH:

        text = text[
            :MAX_TEXT_LENGTH
        ]

    return text.strip()


# =========================================================
# EXTRACT PDF TEXT
# =========================================================

def extract_pdf_text(
    file_bytes: bytes
) -> str:
    """
    Extract text from a PDF file.

        Uses pypdf (or the compatible PyPDF2 package).
    """

    if not PDF_AVAILABLE:

        raise ImportError(
            "pypdf is not installed. "
            "Install it using: pip install pypdf"
        )

    if not file_bytes:

        return ""

    text_parts = []

    pdf_stream = io.BytesIO(
        file_bytes
    )

    reader = PyPDF2.PdfReader(
        pdf_stream
    )

    for page_number, page in enumerate(
        reader.pages,
        start=1
    ):

        try:

            page_text = page.extract_text()

        except Exception as error:

            page_text = (
                f"\n[Unable to extract page "
                f"{page_number}: {error}]\n"
            )

        if page_text:

            text_parts.append(
                page_text
            )

    return "\n".join(
        text_parts
    )


# =========================================================
# EXTRACT DOCX TEXT
# =========================================================

def extract_docx_text(
    file_bytes: bytes
) -> str:
    """
    Extract text from a DOCX file.

    Reads:
        - Paragraphs
        - Tables
    """

    if not DOCX_AVAILABLE:

        raise ImportError(
            "python-docx is not installed. "
            "Install it using: pip install python-docx"
        )

    if not file_bytes:

        return ""

    document_stream = io.BytesIO(
        file_bytes
    )

    document = Document(
        document_stream
    )

    text_parts = []

    # -----------------------------------------------------
    # Paragraphs
    # -----------------------------------------------------

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:

            text_parts.append(
                text
            )

    # -----------------------------------------------------
    # Tables
    # -----------------------------------------------------

    for table in document.tables:

        for row in table.rows:

            cells = []

            for cell in row.cells:

                cell_text = cell.text.strip()

                if cell_text:

                    cells.append(
                        cell_text
                    )

            if cells:

                text_parts.append(
                    " | ".join(cells)
                )

    return "\n".join(
        text_parts
    )


# =========================================================
# EXTRACT TXT TEXT
# =========================================================

def extract_txt_text(
    file_bytes: bytes
) -> str:
    """
    Extract text from a TXT file.

    Attempts UTF-8 first and then common
    fallback encodings.
    """

    if not file_bytes:

        return ""

    encodings = [
        "utf-8",
        "utf-8-sig",
        "cp1252",
        "latin-1"
    ]

    for encoding in encodings:

        try:

            return file_bytes.decode(
                encoding
            )

        except UnicodeDecodeError:

            continue

    return file_bytes.decode(
        "utf-8",
        errors="ignore"
    )


# =========================================================
# GENERIC FILE EXTRACTION
# =========================================================

def extract_text_from_bytes(
    file_name: str,
    file_bytes: bytes
) -> str:
    """
    Extract text based on file extension.
    """

    validation = validate_file(
        file_name
    )

    if not validation["valid"]:

        raise ValueError(
            validation["message"]
        )

    extension = validation[
        "extension"
    ]

    if extension == ".pdf":

        text = extract_pdf_text(
            file_bytes
        )

    elif extension == ".docx":

        text = extract_docx_text(
            file_bytes
        )

    elif extension == ".txt":

        text = extract_txt_text(
            file_bytes
        )

    else:

        raise ValueError(
            "Unsupported file format."
        )

    return clean_text(
        text
    )


# =========================================================
# STREAMLIT UPLOADED FILE
# =========================================================

def extract_text_from_uploaded_file(
    uploaded_file
) -> str:
    """
    Extract text from a Streamlit UploadedFile.

    Example:

        uploaded_file = st.file_uploader(...)

        text = extract_text_from_uploaded_file(
            uploaded_file
        )
    """

    if uploaded_file is None:

        raise ValueError(
            "No file was uploaded."
        )

    file_name = uploaded_file.name

    file_bytes = uploaded_file.getvalue()

    return extract_text_from_bytes(
        file_name,
        file_bytes
    )


# =========================================================
# RESUME METADATA
# =========================================================

def get_resume_metadata(
    file_name: str,
    text: str
) -> Dict[str, Any]:
    """
    Generate basic metadata from extracted resume text.
    """

    extension = Path(
        file_name
    ).suffix.lower()

    word_count = len(
        re.findall(
            r"\b\w+\b",
            text
        )
    )

    character_count = len(
        text
    )

    line_count = len(
        [
            line
            for line in text.splitlines()
            if line.strip()
        ]
    )

    return {
        "file_name": file_name,
        "file_type": extension,
        "character_count": character_count,
        "word_count": word_count,
        "line_count": line_count,
        "has_text": bool(
            text.strip()
        )
    }


# =========================================================
# PARSE RESUME
# =========================================================

def parse_resume(
    file_name: str,
    file_bytes: bytes
) -> Dict[str, Any]:
    """
    Complete resume parsing pipeline.

    Returns:

        {
            "success": True,
            "file_name": "...",
            "text": "...",
            "metadata": {...}
        }
    """

    try:

        text = extract_text_from_bytes(
            file_name=file_name,
            file_bytes=file_bytes
        )

        if not text.strip():

            return {
                "success": False,
                "file_name": file_name,
                "text": "",
                "metadata": {},
                "error": (
                    "No readable text was found "
                    "in the uploaded resume."
                )
            }

        metadata = get_resume_metadata(
            file_name=file_name,
            text=text
        )

        return {
            "success": True,
            "file_name": file_name,
            "text": text,
            "metadata": metadata,
            "error": None
        }

    except Exception as error:

        return {
            "success": False,
            "file_name": file_name,
            "text": "",
            "metadata": {},
            "error": str(error)
        }


def process_resume(
    file,
    filename: Optional[str] = None
) -> Dict[str, Any]:
    """Parse a Streamlit uploaded file or a file-like object."""

    file_name = filename or getattr(file, "name", None)

    validation = validate_file(file_name)
    if not validation["valid"]:
        return {
            "success": False,
            "file_name": file_name,
            "text": "",
            "metadata": {},
            "validation": validation,
            "error": validation["message"]
        }

    if hasattr(file, "getvalue"):
        file_bytes = file.getvalue()
    elif hasattr(file, "read"):
        file_bytes = file.read()
    elif isinstance(file, (bytes, bytearray)):
        file_bytes = bytes(file)
    else:
        return {
            "success": False,
            "file_name": file_name,
            "text": "",
            "metadata": {},
            "validation": validation,
            "error": "Uploaded file content could not be read."
        }

    result = parse_resume(
        file_name=file_name,
        file_bytes=file_bytes
    )
    result["validation"] = validation
    return result


def detect_resume_sections(resume_text: str) -> Dict[str, str]:
    """Compatibility name for the resume section extractor."""

    return extract_sections(resume_text)


def create_resume_preview(
    resume_text: str,
    max_characters: int = 1200
) -> str:
    """Return a readable, length-limited resume preview."""

    preview = (resume_text or "").strip()
    if max_characters <= 0:
        return ""
    if len(preview) <= max_characters:
        return preview
    return preview[:max_characters].rstrip() + "…"


# =========================================================
# EXTRACT RESUME SECTIONS
# =========================================================

def extract_sections(
    resume_text: str
) -> Dict[str, str]:
    """
    Attempt to identify common resume sections.

    This is a lightweight rule-based parser.
    The Ollama agent will perform deeper semantic analysis.
    """

    if not resume_text:

        return {}

    lines = [
        line.strip()
        for line in resume_text.splitlines()
        if line.strip()
    ]

    section_names = {

        "summary": [
            "summary",
            "professional summary",
            "profile",
            "objective",
            "career objective"
        ],

        "skills": [
            "skills",
            "technical skills",
            "core skills",
            "key skills",
            "technical competencies"
        ],

        "experience": [
            "experience",
            "work experience",
            "professional experience",
            "employment history"
        ],

        "education": [
            "education",
            "academic background",
            "qualifications"
        ],

        "projects": [
            "projects",
            "personal projects",
            "academic projects"
        ],

        "certifications": [
            "certifications",
            "certificates",
            "licenses"
        ],

        "achievements": [
            "achievements",
            "awards",
            "honors"
        ]
    }

    normalized_sections = {}

    for section, names in section_names.items():

        for name in names:

            normalized_sections[
                name
            ] = section

    sections = {}

    current_section = "general"

    sections[current_section] = []

    for line in lines:

        normalized_line = (
            re.sub(
                r"[^a-zA-Z ]",
                "",
                line.lower()
            )
            .strip()
        )

        if normalized_line in normalized_sections:

            current_section = normalized_sections[
                normalized_line
            ]

            if current_section not in sections:

                sections[current_section] = []

            continue

        sections[
            current_section
        ].append(
            line
        )

    # Convert lists to strings

    result = {}

    for section, values in sections.items():

        value = "\n".join(
            values
        ).strip()

        if value:

            result[section] = value

    return result


# =========================================================
# STREAMLIT RESUME UPLOADER
# =========================================================

def resume_uploader(
    label: str = "Upload your resume"
):
    """
    Create a Streamlit resume uploader.

    Returns:
        Streamlit UploadedFile or None.
    """

    import streamlit as st

    uploaded_file = st.file_uploader(
        label,
        type=[
            "pdf",
            "docx",
            "txt"
        ],
        help=(
            "Upload your resume in PDF, "
            "DOCX or TXT format."
        )
    )

    return uploaded_file


# =========================================================
# DISPLAY PARSING RESULT
# =========================================================

def display_resume_preview(
    parsed_resume: Dict[str, Any]
):
    """
    Display extracted resume information
    inside Streamlit.
    """

    import streamlit as st

    if not parsed_resume.get(
        "success"
    ):

        st.error(
            parsed_resume.get(
                "error",
                "Unable to parse resume."
            )
        )

        return

    metadata = parsed_resume.get(
        "metadata",
        {}
    )

    st.success(
        "Resume successfully parsed."
    )

    # -----------------------------------------------------
    # Metadata
    # -----------------------------------------------------

    col1, col2, col3 = st.columns(3)

    with col1:

        st.metric(
            "Words",
            metadata.get(
                "word_count",
                0
            )
        )

    with col2:

        st.metric(
            "Characters",
            metadata.get(
                "character_count",
                0
            )
        )

    with col3:

        st.metric(
            "Lines",
            metadata.get(
                "line_count",
                0
            )
        )

    # -----------------------------------------------------
    # Preview
    # -----------------------------------------------------

    with st.expander(
        "📄 View Extracted Resume Text"
    ):

        st.text_area(
            "Resume Content",
            parsed_resume.get(
                "text",
                ""
            ),
            height=400,
            disabled=True
        )


# =========================================================
# TEST
# =========================================================

if __name__ == "__main__":

    print("=" * 70)

    print(
        "AI RESUME & SKILL GAP AGENT"
    )

    print(
        "Resume Parser Test"
    )

    print("=" * 70)

    # -----------------------------------------------------
    # TXT test
    # -----------------------------------------------------

    sample_text = """
    John Doe

    Python Developer

    SUMMARY
    Python developer with experience in data analysis.

    SKILLS
    Python, SQL, Pandas, Streamlit, Machine Learning

    EXPERIENCE
    Software Developer at ABC Company.

    EDUCATION
    Bachelor of Computer Applications
    """

    sample_bytes = sample_text.encode(
        "utf-8"
    )

    result = parse_resume(
        file_name="sample.txt",
        file_bytes=sample_bytes
    )

    if result["success"]:

        print(
            "\n✅ Resume parsed successfully."
        )

        print(
            "\nMetadata:"
        )

        for key, value in result[
            "metadata"
        ].items():

            print(
                f"{key}: {value}"
            )

        print(
            "\nExtracted sections:"
        )

        sections = extract_sections(
            result["text"]
        )

        for section, content in sections.items():

            print(
                f"\n--- {section.upper()} ---"
            )

            print(
                content
            )

    else:

        print(
            "\n❌ Resume parsing failed."
        )

        print(
            result["error"]
        )

    print(
        "\n" + "=" * 70
    )

